"""Render seed/corpus_manifest.yaml into the files under seed/corpus/.

PDFs go through WeasyPrint because its text layer round-trips Arabic: characters come back
in logical order per glyph run, which api/app/ingest/pdf.py then reorders into reading
order. PyMuPDF's own HTML writer and headless Chrome both produced unrecoverable glyph
maps for Amiri and Noto Naskh in testing.

The rendered files are committed, so nobody needs WeasyPrint (or pango) to run the demo.
Regenerate with:  make corpus
"""

import datetime as dt
import html
import shutil
from pathlib import Path
from typing import Any

import docx
import openpyxl
import yaml
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from hijridate import Gregorian
from openpyxl.styles import Alignment, Font
from weasyprint import CSS as WeasyCSS
from weasyprint import HTML
from weasyprint.text.fonts import FontConfiguration

SEED = Path(__file__).parent
CORPUS = SEED / "corpus"
FONTS = SEED / "fonts"

AR_MONTHS = [
    "يناير",
    "فبراير",
    "مارس",
    "أبريل",
    "مايو",
    "يونيو",
    "يوليو",
    "أغسطس",
    "سبتمبر",
    "أكتوبر",
    "نوفمبر",
    "ديسمبر",
]
AR_DIGITS = str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩")


def arabic_date(d: dt.date) -> str:
    h = Gregorian(d.year, d.month, d.day).to_hijri()
    greg = f"{d.day} {AR_MONTHS[d.month - 1]} {d.year}م".translate(AR_DIGITS)
    hijri = f"{h.day} {h.month_name('ar')} {h.year}هـ".translate(AR_DIGITS)
    return f"{hijri} الموافق {greg}"


def english_date(d: dt.date) -> str:
    h = Gregorian(d.year, d.month, d.day).to_hijri()
    return f"{d.day} {d.strftime('%B %Y')} ({h.day} {h.month_name('en')} {h.year} AH)"


CSS = f"""
@font-face {{ font-family: Amiri; src: url("{(FONTS / "Amiri-Regular.ttf").as_uri()}"); }}
@font-face {{ font-family: Amiri; font-weight: bold; src: url("{(FONTS / "Amiri-Bold.ttf").as_uri()}"); }}
@page {{ size: A4; margin: 22mm 20mm; }}
body {{ font-family: Amiri, serif; font-size: 12pt; line-height: 1.55; color: #1d1d1b;
        font-variant-ligatures: none; }}
.head {{ border-bottom: 1.5pt solid #1d1d1b; padding-bottom: 6pt; margin-bottom: 14pt; }}
.org {{ font-size: 15pt; font-weight: bold; }}
.org-alt {{ font-size: 12pt; }}
.meta {{ margin: 0 0 14pt 0; }}
.meta div {{ margin: 0; }}
h1 {{ font-size: 15pt; margin: 0 0 12pt 0; }}
h2 {{ font-size: 12.5pt; margin: 14pt 0 4pt 0; }}
p {{ margin: 0 0 8pt 0; }}
.sig {{ margin-top: 26pt; }}
.stamp {{ margin-top: 20pt; font-size: 10pt; color: #7a1f1f; font-weight: bold; }}
"""


def org_names(orgs: dict[str, Any], org_id: str | None) -> tuple[str, str]:
    if not org_id:
        return "Linearyze", "لينيرايز"
    o = orgs[org_id]
    return o["name_en"], o["name_ar"]


def render_pdf(doc: dict[str, Any], orgs: dict[str, Any], out: Path) -> None:
    rtl = doc["language"] == "ar"
    e = html.escape
    sender_en, sender_ar = org_names(orgs, doc.get("sender"))
    parts = [f'<html lang="{doc["language"]}" dir="{"rtl" if rtl else "ltr"}"><body>']

    if doc["format"] == "letter":
        primary, secondary = (sender_ar, sender_en) if rtl else (sender_en, sender_ar)
        parts.append(
            f'<div class="head"><div class="org">{e(primary)}</div>'
            f'<div class="org-alt">{e(secondary)}</div></div>'
        )
        date = doc["document_date"]
        if rtl:
            parts.append(
                f'<div class="meta"><div>الرقم: {e(doc["reference_no"])}</div>'
                f"<div>التاريخ: {e(arabic_date(date))}</div></div>"
            )
        else:
            parts.append(
                f'<div class="meta"><div>Ref: {e(doc["reference_no"])}</div>'
                f"<div>Date: {e(english_date(date))}</div></div>"
            )
        subject_label = "الموضوع" if rtl else "Subject"
        parts.append(f"<h1>{subject_label}: {e(doc['title'])}</h1>")
    else:
        parts.append('<div class="head"><div class="org">Linearyze | لينيرايز</div></div>')
        parts.append(f"<h1>{e(doc['title'])}</h1>")
        if doc.get("contract_no"):
            parts.append(
                f'<div class="meta"><div>Contract no: {e(doc["contract_no"])}</div>'
                f"<div>Date: {e(doc['document_date'].strftime('%d %B %Y'))}</div></div>"
            )
        elif not rtl:
            parts.append(
                f'<div class="meta"><div>Date: '
                f"{e(doc['document_date'].strftime('%d %B %Y'))}</div></div>"
            )
        else:
            parts.append(
                f'<div class="meta"><div>التاريخ: {e(arabic_date(doc["document_date"]))}'
                f"</div></div>"
            )

    for para in doc.get("body", []):
        parts.append(f"<p>{e(para)}</p>")
    for section in doc.get("sections", []):
        parts.append(f"<h2>{e(section['heading'])}</h2>")
        parts.extend(f"<p>{e(p)}</p>" for p in section["paragraphs"])

    if doc.get("signatory"):
        parts.append(
            f'<div class="sig"><div>{e(doc["signatory"])}</div>'
            f"<div>{e(doc['signatory_title'])}</div></div>"
        )
    if doc["confidentiality"] == "confidential":
        parts.append(f'<div class="stamp">{"سري" if rtl else "CONFIDENTIAL"}</div>')
    parts.append("</body></html>")

    fonts = FontConfiguration()
    stylesheet = WeasyCSS(string=CSS, font_config=fonts)
    HTML(string="".join(parts)).write_pdf(out, stylesheets=[stylesheet], font_config=fonts)


def _set_rtl(paragraph: Any) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    ppr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def render_docx(doc: dict[str, Any], out: Path) -> None:
    rtl = doc["language"] == "ar"
    d = docx.Document()
    d.core_properties.title = doc["title"]
    d.core_properties.author = "Linearyze"
    paragraphs = [d.add_heading(doc["title"], level=0)]
    for section in doc["sections"]:
        paragraphs.append(d.add_heading(section["heading"], level=1))
        paragraphs.extend(d.add_paragraph(p) for p in section["paragraphs"])
    if rtl:
        for p in paragraphs:
            _set_rtl(p)
    d.save(out)


def render_xlsx(doc: dict[str, Any], out: Path) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.properties.title = doc["title"]
    wb.properties.creator = "Linearyze"
    for sheet in doc["sheets"]:
        ws = wb.create_sheet(sheet["name"])
        for r, row in enumerate(sheet["rows"], start=1):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=value)
                if isinstance(value, int):
                    cell.number_format = "#,##0"
                    cell.alignment = Alignment(horizontal="right")
                if r == 1:
                    cell.font = Font(bold=True, size=13)
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = 34
    wb.save(out)


def main() -> None:
    manifest = yaml.safe_load((SEED / "corpus_manifest.yaml").read_text())
    orgs = {o["id"]: o for o in yaml.safe_load((SEED / "users.yaml").read_text())["organizations"]}
    if CORPUS.exists():
        shutil.rmtree(CORPUS)
    for doc in manifest["documents"]:
        out = CORPUS / doc["file"]
        out.parent.mkdir(parents=True, exist_ok=True)
        if doc["format"] == "docx":
            render_docx(doc, out)
        elif doc["format"] == "xlsx":
            render_xlsx(doc, out)
        else:
            render_pdf(doc, orgs, out)
    print(f"wrote {len(manifest['documents'])} documents to {CORPUS}")  # noqa: T201


if __name__ == "__main__":
    main()
